import sys,unittest,io
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from analysis import analyze,fit_sample,four_pl
from agent import run_workflow
from reports import make_docx
from sample_support import rewrite_wells
ROOT=Path(__file__).resolve().parents[2]
DATA=next(p for p in ROOT.iterdir() if p.is_dir() and (p/"RAW DATA.csv").exists())
RAW=(DATA/"RAW DATA.csv").read_bytes()
REC=next(DATA.glob("*.xlsx")).read_bytes()
class AssayTests(unittest.TestCase):
    def test_synthetic_data_recovers_known_generating_parameters(self):
        r=analyze(RAW,REC)
        for s,expected in zip(r["samples"],[3.2,7.5,18.0]):
            self.assertAlmostEqual(s["ic50_nm"],expected,delta=.03)
        self.assertIsNone(r["samples"][3]["ic50_nm"])
        self.assertGreater(r["controls"]["z_prime"],.9)
        self.assertEqual(len(r["provenance"]["design"]),72)
    def test_changed_data_changes_estimate(self):
        raw=rewrite_wells(RAW,{"E3":100000,"E4":100000})
        self.assertGreater(abs(analyze(raw,REC)["samples"][0]["ic50_nm"]-analyze(RAW,REC)["samples"][0]["ic50_nm"]),.1)
    def test_result_sheet_is_never_consumed(self):
        wb=load_workbook(io.BytesIO(REC));wb["Result"]["C11"]=999999
        out=io.BytesIO();wb.save(out)
        self.assertAlmostEqual(analyze(RAW,out.getvalue())["samples"][0]["ic50_nm"],analyze(RAW,REC)["samples"][0]["ic50_nm"],places=8)
    def test_working_concentration_is_not_final_concentration(self):
        wb=load_workbook(io.BytesIO(REC));wb["protocol"]["D31"]=999999
        out=io.BytesIO();wb.save(out)
        self.assertAlmostEqual(analyze(RAW,out.getvalue())["samples"][0]["ic50_nm"],3.2,delta=.03)
    def test_known_synthetic_curve(self):
        x=np.repeat(np.geomspace(.1,100,12),2)
        y=four_pl(x,10000,300000,np.log10(3.2),1.3)
        self.assertAlmostEqual(fit_sample(x,y,310000,10000)["ic50_nm"],3.2,places=4)
    def test_flat_curve_is_not_forced(self):
        r=fit_sample(np.repeat(np.geomspace(.1,100,8),2),np.ones(16)*300000,310000,10000)
        self.assertIsNone(r["ic50_nm"]);self.assertEqual(r["curve"],[])
    def test_missing_well_is_rejected(self):
        with self.assertRaisesRegex(ValueError,"缺少原始数据"):analyze(rewrite_wells(RAW,{"A3":""}),REC)
    def test_invalid_signal_is_rejected(self):
        with self.assertRaises(ValueError):analyze(rewrite_wells(RAW,{"A3":"NaN"}),REC)
    def test_instrument_background_is_recorded_not_subtracted(self):
        result=analyze(RAW,REC)
        self.assertEqual(result["controls"]["instrument_background"],0)
        self.assertTrue(any("Background information" in a["source"] for a in result["alerts"]))
    def test_reference_control_is_explicit_and_never_inferred(self):
        base=analyze(RAW,REC)
        self.assertIsNone(base["controls"]["reference_control"]["sample"])
        self.assertTrue(any("未指定药效参考内控" in a["text"] for a in base["alerts"]))
        result=analyze(RAW,REC,{"reference_sample":"Drug1","reference_ic50":2.4})
        self.assertEqual(result["controls"]["reference_control"]["sample"],"Drug1")
        self.assertAlmostEqual(result["controls"]["reference_control"]["ratio_to_expected"],result["samples"][0]["ic50_nm"]/2.4,places=6)
    def test_reference_control_mismatch_and_unknown_identity(self):
        result=analyze(RAW,REC,{"reference_sample":"Drug1","reference_ic50":20})
        self.assertIn("需复核",result["controls"]["reference_control"]["status"])
        self.assertIsNotNone(result["samples"][0]["ic50_nm"])
        self.assertIn("参考内控未通过",result["samples"][0]["status"])
        with self.assertRaises(ValueError):analyze(RAW,REC,{"reference_sample":"NotInDesign"})
        with self.assertRaises(ValueError):analyze(RAW,REC,{"reference_sample":"Drug1","reference_ic50":-1})
    def test_blank_control_qc_flags_disagreement_and_keeps_raw_wells(self):
        wb=load_workbook(io.BytesIO(REC));ws=wb["protocol"]
        ws["O44"]="Blank";ws["O45"]="Blank"
        record=io.BytesIO();wb.save(record)
        rows=list(__import__("csv").reader(io.StringIO(RAW.decode("utf-8"))))
        for row in rows:
            if row and row[0]=="A":row[12]="100"
            if row and row[0]=="B":row[12]="1000"
        raw=io.StringIO();__import__("csv").writer(raw).writerows(rows)
        result=analyze(raw.getvalue().encode(),record.getvalue())
        self.assertEqual(result["controls"]["blank_check"]["count"],2)
        self.assertEqual(result["controls"]["blank_check"]["status"],"需复核")
        self.assertEqual(result["provenance"]["raw_wells"]["B12"],1000)
    def test_abnormal_replicate_is_flagged_but_preserved(self):
        result=analyze(rewrite_wells(RAW,{"A3":200000}),REC)
        self.assertEqual(result["provenance"]["raw_wells"]["A3"],200000)
        self.assertTrue(any("A3" in a["source"] and a["level"]=="warning" for a in result["alerts"]))
        self.assertEqual(sum(len(p["wells"]) for p in result["samples"][0]["points"]),16)
    def test_reversed_controls_are_rejected(self):
        wb=load_workbook(io.BytesIO(REC))
        for row in range(44,52):wb["protocol"][f"N{row}"]="Cell only" if row<48 else "No Drug"
        out=io.BytesIO();wb.save(out)
        with self.assertRaisesRegex(ValueError,"No Drug"):analyze(RAW,out.getvalue())
    def test_quality_threshold_validation(self):
        with self.assertRaises(ValueError):analyze(RAW,REC,{"cv_limit":0})
    def test_non_nm_design_is_rejected(self):
        wb=load_workbook(io.BytesIO(REC));wb["protocol"]["C41"]="Plate Map concentration (uM)"
        out=io.BytesIO();wb.save(out)
        with self.assertRaisesRegex(ValueError,"nM"):analyze(RAW,out.getvalue())
    def test_zero_baseline_cv_does_not_crash(self):
        import csv
        rows=list(csv.reader(io.StringIO(RAW.decode("utf-8"))))
        for row in rows:
            if row and row[0] in ("E","F","G","H"):row[11]="0"
        out=io.StringIO();csv.writer(out).writerows(rows)
        result=analyze(out.getvalue().encode(),REC)
        self.assertIsNone(result["controls"]["cell_only_cv"])
        self.assertGreater(result["controls"]["z_prime"],.9)
    def test_workflow_and_report(self):
        b={"raw":RAW,"record":REC,"raw_name":"RAW DATA.csv","record_name":"record.xlsx","name":"合成验收批次"}
        r=run_workflow([b],{},lambda *args:None)
        doc=make_docx(r);self.assertTrue(doc.startswith(b"PK"))
        from docx import Document
        d=Document(io.BytesIO(doc));text="\n".join(p.text for p in d.paragraphs)
        self.assertIn("未提供",text);self.assertIn("附件、溯源",text);self.assertIn("SHA256",text)
        self.assertEqual(len(d.sections[0].header.tables),1)
        self.assertTrue(any("AssayPilot" in c.text for row in d.sections[0].header.tables[0].rows for c in row.cells))
if __name__=="__main__":unittest.main()
