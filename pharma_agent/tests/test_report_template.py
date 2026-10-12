"""Verify the self-authored public report without an enterprise template."""
import io,unittest,zipfile,sys
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from analysis import analyze
from reports import make_docx,make_brief_docx,make_html
DATA=Path(__file__).resolve().parents[2]/'data'
RESULT={'results':[analyze((DATA/'RAW DATA.csv').read_bytes(),(DATA/'record.xlsx').read_bytes(),batch='合成演示')],'mode':'本地规则编排（未调用大模型）','trace':[]}
class PublicReportTests(unittest.TestCase):
    def test_repeating_header_is_self_authored_and_contains_page_field(self):
        doc=Document(io.BytesIO(make_docx(RESULT)))
        self.assertEqual(len(doc.sections[0].header.tables),1)
        content=' '.join(c.text for row in doc.sections[0].header.tables[0].rows for c in row.cells)
        self.assertIn('AssayPilot',content);self.assertIn('SYNTHETIC DEMO',content)
        self.assertTrue(any(e.get(qn('w:instr'))=='PAGE' for e in doc.sections[0].header._element.iter(qn('w:fldSimple'))))
    def test_result_table_contains_evidence_not_fake_result_sheet(self):
        doc=Document(io.BytesIO(make_docx(RESULT)))
        text=' '.join(c.text for table in doc.tables for row in table.rows for c in row.cells)
        self.assertIn('Drug4',text);self.assertIn('不可估计',text);self.assertNotIn('999999',text);self.assertIn('3.2',text)
    def test_no_embedded_office_object_or_external_template(self):
        raw=make_docx(RESULT)
        with zipfile.ZipFile(io.BytesIO(raw)) as z:self.assertFalse(any(n.startswith('word/embeddings/') for n in z.namelist()))
        self.assertEqual(len(list(Document(io.BytesIO(raw)).element.body.iter(qn('a:blip')))),1)
    def test_human_edit_does_not_modify_numeric_results(self):
        doc=Document(io.BytesIO(make_docx(RESULT,edited_summary='人工复核：Drug4需确认。')))
        self.assertIn('人工复核：Drug4需确认。',[p.text for p in doc.paragraphs]);self.assertAlmostEqual(RESULT['results'][0]['samples'][0]['ic50_nm'],3.2,delta=.03)
    def test_brief_contains_nine_sections_and_missing_controls(self):
        doc=Document(io.BytesIO(make_brief_docx(RESULT)));text='\n'.join(p.text for p in doc.paragraphs)
        for n in range(1,10):self.assertIn(str(n)+'. ',text)
        self.assertIn('独立Blank',text);self.assertIn('参考内控',text);self.assertIn('EC50',' '.join(c.text for t in doc.tables for row in t.rows for c in row.cells))
    def test_html_escapes_human_text(self):
        raw=make_html(RESULT,edited_summary='<script>alert(1)</script>').decode()
        self.assertNotIn('<script>alert(1)</script>',raw);self.assertIn('&lt;script&gt;',raw)
if __name__=='__main__':unittest.main()
