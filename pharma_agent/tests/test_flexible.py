import copy,csv,io,json,math,sys,unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
from docx import Document
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'pharma_agent'))
from analysis import analyze
from imports import read_tables,long_table,describe
from datasets import preview
from reports import make_docx,make_html,summary
import agent,server
DATA=ROOT/'data'/'flexible'
def data_rows(direction='inhibition'):
    out=[]
    for c in [.03,.1,.3,1,3,10,30,100,300]:
        signal=100+900/(1+(c/5)**1.4)
        if direction=='activation':signal=1100-signal
        for f in [.995,1.005]:out.append(dict(sample='Independent',concentration=c,unit='nM',value=signal*f,role='sample'))
    for role,val in [('high',1000),('low',100)]:
        for f in [.995,1.005]:out.append(dict(sample=role,concentration='',unit='nM',value=val*f,role=role))
    return out
def json_data(rows):return json.dumps(rows).encode()
class FlexibleTests(unittest.TestCase):
    def test_csv_tsv_semicolon_text_json_and_xlsx_agree(self):
        for file,scene in [('inhibition.csv','inhibition'),('inhibition.tsv','inhibition'),('inhibition.txt','inhibition'),('inhibition.json','inhibition'),('activation.xlsx','activation')]:
            with self.subTest(file=file):
                r=analyze((DATA/file).read_bytes(),options={'assay_type':scene},raw_name=file)
                self.assertAlmostEqual(r['samples'][0]['potency_nm'],5.,places=5)
                self.assertEqual(r['provenance']['import']['conversions'],['uM→nM'])
    def test_ascending_curve_is_ec50_not_ic50(self):
        r=analyze(json_data(data_rows('activation')),options={'assay_type':'activation','readout':'RFU'},raw_name='a.json')
        s=r['samples'][0];self.assertAlmostEqual(s['ec50_nm'],5.,places=5);self.assertIsNone(s['ic50_nm'])
        self.assertLess(s['curve'][0]['rlu'],s['curve'][-1]['rlu']);self.assertLess(s['curve'][0]['effect_pct'],s['curve'][-1]['effect_pct'])
        self.assertAlmostEqual(s['absolute50_nm'],5.,places=5)
    def test_viability_and_enzyme_have_correct_semantics(self):
        for name,field in [('viability','cc50_nm'),('enzyme','ic50_nm')]:
            r=analyze(json_data(data_rows()),options={'assay_type':name,'readout':'OD'},raw_name='a.json');self.assertAlmostEqual(r['samples'][0][field],5.,places=5)
            if name=='viability':self.assertIsNone(r['samples'][0]['ic50_nm']);self.assertIn('不等于',r['assay']['principle'])
    def test_wrong_direction_and_flat_curves_are_not_forced(self):
        r=analyze(json_data(data_rows()),options={'assay_type':'activation'},raw_name='a.json');self.assertIsNone(r['samples'][0]['ec50_nm'])
        rows=data_rows();
        for row in rows:
            if row['role']=='sample':row['value']=700
        self.assertIsNone(analyze(json_data(rows),raw_name='a.json')['samples'][0]['potency_nm'])
    def test_custom_header_selection_and_mapping(self):
        raw=(DATA/'custom-headers.xlsx').read_bytes()
        with self.assertRaises(ValueError):analyze(raw,raw_name='x.xlsx')
        config={'raw':{'sheet':'Measurements','header_row':2,'columns':dict(zip(['sample','concentration','unit','value','role'],['compound_code','dose_level','dose_units','sensor_reading','group_type']))}}
        r=analyze(raw,raw_name='x.xlsx',import_config=config)
        self.assertAlmostEqual(r['samples'][0]['potency_nm'],5.,places=5)
        self.assertEqual(r['provenance']['import']['mapping']['value'],'sensor_reading')
    def test_384_matrix_joins_tsv_by_well(self):
        r=analyze((DATA/'matrix-384.xlsx').read_bytes(),(DATA/'design-384.tsv').read_bytes(),raw_name='plate.xlsx',record_name='map.tsv')
        self.assertAlmostEqual(r['samples'][0]['ic50_nm'],5.,places=5)
        self.assertTrue(any(d['well'].startswith('K') for d in r['metadata']['design']))
    def test_chinese_headers_unicode_units_and_column_order(self):
        rows=data_rows();headers=['读数','孔类型','单位','终浓度','样品名称'];out=io.StringIO();w=csv.writer(out);w.writerow(headers)
        for row in rows:w.writerow([row['value'],row['role'],'µM',row['concentration']/1000 if row['concentration']!='' else '',row['sample']])
        r=analyze(out.getvalue().encode(),raw_name='中文.csv');self.assertAlmostEqual(r['samples'][0]['ic50_nm'],5.,places=5)
    def test_mass_units_nonfinite_negative_and_unknown_roles_rejected(self):
        for field,value in [('unit','mg/mL'),('value','NaN'),('value',-10),('concentration',-1),('role','positive')]:
            with self.subTest(field=field,value=value):
                rows=data_rows();rows[0][field]=value
                with self.assertRaises(ValueError):analyze(json_data(rows),raw_name='a.json')
    def test_missing_controls_are_explicit(self):
        rows=[r for r in data_rows() if r['role']=='sample']
        with self.assertRaisesRegex(ValueError,'至少各需2'):analyze(json_data(rows),raw_name='a.json')
    def test_header_and_unit_column_conflict_is_rejected(self):
        rows=[{'sample':'A','concentration (nM)':1,'unit':'uM','value':100,'role':'sample'}]
        with self.assertRaisesRegex(ValueError,'单位.*冲突'):analyze(json_data(rows),raw_name='a.json')
    def test_subnanomolar_readout_can_be_fitted_without_rlu_assumption(self):
        rows=data_rows()
        for row in rows:row['value']/=1000
        r=analyze(json_data(rows),options={'assay_type':'enzyme','readout':'OD'},raw_name='od.json')
        self.assertAlmostEqual(r['samples'][0]['ic50_nm'],5.,places=5)
    def test_duplicate_physical_wells_are_rejected(self):
        rows=data_rows();rows[0]['well']='A1';rows[1]['well']='A1'
        with self.assertRaisesRegex(ValueError,'重复'):analyze(json_data(rows),raw_name='a.json')
    def test_zero_concentration_observation_is_retained_not_log_fitted(self):
        rows=data_rows();rows.insert(0,dict(sample='Independent',concentration=0,unit='nM',value=1000,role='sample'))
        r=analyze(json_data(rows),raw_name='a.json');self.assertEqual(r['samples'][0]['points'][0]['concentration'],0);self.assertAlmostEqual(r['samples'][0]['ic50_nm'],5.,places=5)
        self.assertTrue(all(p['concentration']>0 for p in r['samples'][0]['curve']))
    def test_result_sheet_excluded_and_multi_sheet_requires_selection(self):
        wb=Workbook();ws=wb.active;ws.title='first';ws.append(['sample','concentration','value','role']);ws.append(['A',1,500,'sample'])
        clone=wb.copy_worksheet(ws);clone.title='second';wb.create_sheet('Result');out=io.BytesIO();wb.save(out)
        self.assertNotIn('Result',read_tables(out.getvalue(),'x.xlsx')[0])
        with self.assertRaisesRegex(ValueError,'多个工作表'):long_table(out.getvalue(),'x.xlsx',require_design=True)
    def test_report_chat_and_trace_use_ec50(self):
        r=analyze(json_data(data_rows('activation')),options={'assay_type':'activation'},raw_name='a.json')
        self.assertIn('相对EC50',summary(r));self.assertNotIn('相对IC50',summary(r))
        answer=agent.local_answer([r]);self.assertIn('EC50',answer);self.assertIn('激活率',answer)
        result={'results':[r],'mode':'本地规则','trace':[]};doc=Document(io.BytesIO(make_docx(result)))
        self.assertTrue(any('相对EC50' in c.text for table in doc.tables for row in table.rows for c in row.cells))
        self.assertIn('激活',doc.paragraphs[0].text);self.assertIn('EC50',make_html(result).decode())
    def test_probe_is_bounded_and_exposes_selectable_custom_headers(self):
        p=describe((DATA/'custom-headers.xlsx').read_bytes(),'a.xlsx')
        self.assertEqual(len(p['tables']),1);self.assertEqual(p['tables'][0]['headers'][0]['header_row'],2)
        self.assertIn('sensor_reading',p['tables'][0]['headers'][0]['columns'])
    def test_server_accepts_single_files_and_validates_extensions(self):
        import base64
        raw=json_data(data_rows());body={'batches':[{'raw_name':'a.json','raw':base64.b64encode(raw).decode()}]}
        self.assertIsNone(server.decode_batches(body)[0]['record'])
        body['batches'][0]['raw_name']='a.pdf'
        with self.assertRaises(ValueError):server.decode_batches(body)
    def test_pipeline_supplies_assay_and_canonical_evidence_to_model(self):
        batches=[{'raw':json_data(data_rows('activation')),'raw_name':'a.json','name':'activation'}]
        with patch.object(agent,'config',return_value={'enabled':False}):r=agent.run_workflow(batches,{'assay_type':'activation'},lambda *args:None)
        self.assertEqual(agent.compact(r['results'])[0]['assay']['potency_label'],'EC50')
if __name__=='__main__':unittest.main()
