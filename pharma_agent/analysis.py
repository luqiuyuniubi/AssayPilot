"""Deterministic assay tools. Report.docx and the Result worksheet are never inputs."""
import csv, io, math, hashlib, warnings, re
import numpy as np
from scipy.optimize import curve_fit, OptimizeWarning
from scipy.stats import t
from openpyxl import load_workbook
from assays import profile,effect,potency

RULES = {
 "normalization": "抑制率 = 100 × (No Drug均值 − RLU) / (No Drug均值 − Cell only均值)，不截断",
 "fit": "原始RLU、所有复孔等权重4PL：Bottom+(Top-Bottom)/(1+10^((log10(C)-logIC50)*Hill))",
 "qc": "CV阈值默认20%，Z′默认0.5；复孔CV提示疑似离群，拟合残差标准差超过3倍时提示异常孔；仅标记、不自动删孔；均为演示规则，不是已验证SOP",
 "ic50": "相对IC50：拟合Top与Bottom中点浓度；绝对50%抑制浓度另列；本实验是阻断实验，不计算EC50",
 "uncertainty": "logIC50协方差与t分布近似95%区间；不等同Prism的profile-likelihood区间",
}
def digest(b): return hashlib.sha256(b).hexdigest()
def clean(v):
    if v is None or not str(v).strip() or set(str(v).strip()) == {"*"}: return "未提供"
    return str(v).strip()
def number(v):
    try:
        n=float(v)
        return n if math.isfinite(n) else None
    except (TypeError, ValueError): return None
def parse_raw(data):
    if len(data)>8*1024*1024:raise ValueError('CSV超过8MB限制。')
    text=None
    for enc in ("utf-8-sig","gb18030"):
        try: text=data.decode(enc); break
        except UnicodeDecodeError: pass
    if text is None: raise ValueError("CSV编码无法识别，请使用UTF-8或GB18030。")
    rows=[]
    for row in csv.reader(io.StringIO(text)):
        if len(rows)>=5000:raise ValueError('CSV行数超过安全限制，请按板次拆分文件。')
        if len(row)>256:raise ValueError('CSV列数超过安全限制。')
        rows.append(row)
    wells={}; date="未提供"; header=None
    for row in rows:
        if "Measurement date" in row:
            idx=row.index("Measurement date"); pos=rows.index(row)
            if pos+1<len(rows) and idx<len(rows[pos+1]): date=rows[pos+1][idx]
        if len(row)>=13 and [x.strip() for x in row[1:13]] == [str(i) for i in range(1,13)]:
            header=row; continue
        if header is not None and row and row[0].strip() in list("ABCDEFGH"):
            for col,v in enumerate(row[1:13],1):
                if not v.strip(): continue
                n=number(v)
                if n is None or n<0: raise ValueError(f"{row[0]}{col}的RLU不是有效非负数。")
                key=f"{row[0].strip()}{col}"
                if key in wells: raise ValueError("文件含多块板或重复孔位；请将每块板分别导出并添加为批次。")
                wells[key]=n
    if not wells: raise ValueError("未找到96孔板矩阵：应包含1–12列及A–H行。")
    return wells,date

def parse_instrument_background(data):
    """Read the instrument's recorded background Result without subtracting it."""
    text=None
    for enc in ("utf-8-sig","gb18030"):
        try:text=data.decode(enc);break
        except UnicodeDecodeError:pass
    if text is None:return None
    rows=list(csv.reader(io.StringIO(text)))
    for i,row in enumerate(rows):
        if row and row[0].strip().lower()=="background information":
            for j in range(i+1,min(i+5,len(rows))):
                labels=[v.strip().lower() for v in rows[j]]
                if "plate" in labels and "result" in labels and j+1<len(rows):
                    col=labels.index("result")
                    return number(rows[j+1][col]) if col<len(rows[j+1]) else None
    return None

def parse_record(data):
    if len(data)>8*1024*1024: raise ValueError("实验记录超过8MB限制。")
    import zipfile
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        entries=z.infolist()
        if len(entries)>2000 or sum(i.file_size for i in entries)>40*1024*1024: raise ValueError("实验记录解压体积或条目数过大。")
        from defusedxml.ElementTree import iterparse
        cells=0
        for item in entries:
            if item.file_size>12*1024*1024:raise ValueError('实验记录单个压缩条目过大。')
            if item.filename.endswith(('.xml','.rels')):
                worksheet=item.filename.startswith('xl/worksheets/')
                with z.open(item) as source:
                    for event,element in iterparse(source,events=('start','end'),forbid_dtd=True):
                        if event=='end':element.clear();continue
                        if not worksheet:continue
                        if element.tag.startswith('{') and not element.tag.startswith('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'):continue
                        tag=element.tag.rsplit('}',1)[-1]
                        if tag=='row' and 'r' in element.attrib and (not str(element.get('r','')).isdigit() or not 1<=int(element.get('r'))<=3000):raise ValueError('实验记录工作表行范围超过安全限制。')
                        if tag=='col' and ('min' in element.attrib or 'max' in element.attrib) and any(not str(element.get(k,'')).isdigit() or not 1<=int(element.get(k))<=256 for k in ('min','max')):raise ValueError('实验记录工作表列范围超过安全限制。')
                        if tag=='c':
                            cells+=1
                            if cells>150000:raise ValueError('实验记录单元格数量超过安全限制。')
                        if tag in ('c','dimension','mergeCell'):
                            reference=element.get('r') or element.get('ref') or ''
                            for column,row in re.findall(r'([A-Z]+)(\d+)',reference):
                                col=0
                                for letter in column:col=col*26+ord(letter)-64
                                if not 1<=int(row)<=3000 or not 1<=col<=256:raise ValueError('实验记录工作表范围超过安全限制。')
        if cells==0:raise ValueError('实验记录没有有效的工作表单元格。')
    wb=load_workbook(io.BytesIO(data),data_only=True,read_only=False,keep_links=False)
    ws=next((w for w in wb if w.title.lower()=="protocol"),None)
    if ws is None: raise ValueError("实验记录需包含protocol工作表；不读取Result中的已有结果。")
    anchor=None
    for row in ws:
        for c in row:
            if c.value is not None and "Plate Map" in str(c.value): anchor=c
    if anchor is None: raise ValueError("protocol缺少Plate Map。")
    if not re.search(r"(?<![A-Za-z])nM(?![A-Za-z])",str(anchor.value)):raise ValueError("当前解析器要求Plate Map明确注明nM；其他单位需转换后分析。")
    headrow=anchor.row+2
    rowcol=anchor.column
    colmap={}
    for c in ws[headrow]:
        v=number(c.value)
        if v is not None and v.is_integer() and 1<=v<=12: colmap[c.column]=int(v)
    if len(colmap)!=12: raise ValueError("Plate Map的列编号必须包含1–12。")
    group_at={}
    current=None
    for sheetcol,platecol in sorted(colmap.items()):
        label=ws.cell(headrow-1,sheetcol).value
        if label is not None and str(label).strip(): current=str(label).strip()
        if current: group_at[platecol]=current
    design=[]
    for r in range(headrow+1,headrow+9):
        rowname=str(ws.cell(r,rowcol).value or "").strip()
        if rowname not in "ABCDEFGH" or len(rowname)!=1: raise ValueError("Plate Map缺少A–H行。")
        for sheetcol,platecol in colmap.items():
            v=ws.cell(r,sheetcol).value
            if v is None or str(v).strip()=="": continue
            source=f"protocol!{ws.cell(r,sheetcol).coordinate}"
            if str(v).strip().lower() in ("no drug","cell only","blank"):
                design.append(dict(well=f"{rowname}{platecol}",role=str(v).strip().lower(),source=source))
            else:
                c=number(v)
                if c is None or c<=0: raise ValueError(f"{source}的终浓度需为正数。")
                if platecol not in group_at: raise ValueError(f"{source}没有样品列标题。")
                design.append(dict(well=f"{rowname}{platecol}",role="sample",sample=group_at[platecol],concentration=c,unit="nM",source=source))
    if not design: raise ValueError("Plate Map没有设计孔。")
    instruments=[]; reagents=[]; methods=[]; cell={}; purpose="未提供"; title="报告基因细胞实验"
    section=None
    for row in ws:
        vals=[clean(c.value) for c in row]
        line=" ".join(x for x in vals if x!="未提供")
        if "实验目的" in line: purpose=line
        if "实验名称" in line:
            nonempty=[c.value for c in row if c.value is not None]
            if len(nonempty)>1:title=clean(nonempty[1])
        if "实验试剂：" in line: section="reagents"; continue
        if "实验仪器：" in line: section="instruments"; continue
        if "细胞信息：" in line: section="cell"; continue
        if "实验步骤" in line: section="methods"; continue
        c=clean(ws.cell(row[0].row,3).value)
        if section in ("reagents","instruments") and c!="未提供" and c not in ("试剂名称","仪器名称"):
            entry=[clean(ws.cell(row[0].row,i).value) for i in (3,4,5,6)]
            (reagents if section=="reagents" else instruments).append(entry)
        if section=="cell" and c!="未提供" and c!="细胞":
            cell={"name":c,"source":clean(ws.cell(row[0].row,4).value),"passage":clean(ws.cell(row[0].row,5).value)}
        if section=="methods" and number(ws.cell(row[0].row,2).value) is not None:
            methods.append({"text":c,"source":f"protocol!C{row[0].row}"})
    return dict(title=title,purpose=purpose,instruments=instruments,reagents=reagents,cell=cell,methods=methods,design=design,scope="仅读取protocol；不读取Result或Report.docx中的结论")

def inspect(raw,record):
    wells,date=parse_raw(raw); meta=parse_record(record)
    missing=[d["well"] for d in meta["design"] if d["well"] not in wells]
    if missing: raise ValueError("设计孔缺少原始数据："+", ".join(missing))
    return dict(metadata=meta,date=date,well_count=len(wells),samples=sorted({d["sample"] for d in meta["design"] if d["role"]=="sample"}),raw_sha256=digest(raw),record_sha256=digest(record))

def four_pl(c,bottom,span,log_ic,hill):
    return bottom+span/(1+10**np.clip((np.log10(c)-log_ic)*hill,-100,100))

def fit_sample(concs,values,high,low,mode="inhibition",min_span=20):
    if mode=='activation':
        # Reflect the response about the two control means, fit the same bounded
        # decreasing 4PL, then restore measured signals. EC50 never aliases IC50.
        fit=fit_sample(concs,[high+low-v for v in values],high,low,min_span=min_span)
        fit['ec50_nm']=fit['ic50_nm'];fit['ic50_nm']=None
        if fit['bottom'] is not None:fit['bottom'],fit['top']=high+low-fit['top'],high+low-fit['bottom']
        for point in fit['curve']:point['rlu']=high+low-point['rlu']
        fit['reason']=fit['reason'].replace('IC50','EC50').replace('抑制趋势','激活趋势')
        return fit
    x=np.asarray(concs,dtype=float); y=np.asarray(values,dtype=float)
    unique=sorted(set(x))
    means=np.array([np.mean(y[x==c]) for c in unique])
    norm=100*(high-means)/(high-low)
    amplitude=float(np.ptp(means)/(high-low)*100) if len(means) else 0.
    result={"ic50_nm":None,"ec50_nm":None,"cc50_nm":None,"absolute50_nm":None,"ci95_nm":None,"r2":None,"bottom":None,"top":None,"hill":None,"curve":[],"status":"不可可靠估计","reason":"","dynamic_range_pct":amplitude,"max_inhibition_pct":float(norm[-1]) if len(norm) else None,"max_observed_inhibition_pct":float(max(norm)) if len(norm) else None}
    if len(unique)<5: result["reason"]="少于5个非零浓度点，四参数拟合信息不足";return result
    if amplitude<min_span or means[-1]>=means[0]:
        result["reason"]="信号动态范围不足或不呈抑制趋势；不报告数值IC50";return result
    scale=max(high, float(max(y)),1); yn=y/scale
    lower=[-0.5,0,np.log10(min(x))-3,0.05]; upper=[1.5,2.5,np.log10(max(x))+3,8]
    best=None
    with warnings.catch_warnings():
        warnings.simplefilter("ignore",OptimizeWarning)
        for guess in np.quantile(np.log10(x),[0.25,0.5,0.75]):
            try:
                p,cov=curve_fit(four_pl,x,yn,p0=[max(0,float(min(yn))),max(.01,float(np.ptp(yn))),guess,1],bounds=(lower,upper),maxfev=20000)
                sse=float(np.sum((yn-four_pl(x,*p))**2))
                if best is None or sse<best[0]:best=(sse,p,cov)
            except (RuntimeError,ValueError,FloatingPointError):pass
    if best is None:result["reason"]="拟合未收敛";return result
    sse,p,cov=best; bottom,span,log_ic,hill=map(float,p); ic=10**log_ic
    r2=1-sse/float(np.sum((yn-np.mean(yn))**2))
    result.update(bottom=bottom*scale,top=(bottom+span)*scale,hill=hill,r2=r2)
    result["log_ic50"]=log_ic
    result["residual_sd_rlu"]=float(math.sqrt(sse/(len(x)-4))*scale)
    unreliable=[]
    if not min(x)<=ic<=max(x):unreliable.append("相对IC50超出实测浓度范围")
    if r2<.8:unreliable.append("R²低于0.8")
    if any(abs(p[i]-lower[i])<1e-4 or abs(p[i]-upper[i])<1e-4 for i in range(4)):unreliable.append("参数触及拟合边界")
    if not np.all(np.isfinite(cov)) or np.linalg.cond(cov)>1e12:unreliable.append("参数协方差不稳定")
    ci=None
    if not unreliable:
        se=math.sqrt(max(0,float(cov[2,2]))); delta=float(t.ppf(.975,len(x)-4))*se
        if delta>2:unreliable.append("IC50不确定性过大")
        else:ci=[10**(log_ic-delta),10**(log_ic+delta)]
    grid=np.geomspace(min(x),max(x),120)
    result["curve"]=[{"concentration":float(c),"rlu":float(v),"inhibition_pct":float(100*(high-v)/(high-low))} for c,v in zip(grid,four_pl(grid,*p)*scale)]
    target=(high+low)/2/scale
    if bottom<target<bottom+span:
        absolute=ic*((span/(target-bottom))-1)**(1/hill)
        if min(x)<=absolute<=max(x):result["absolute50_nm"]=float(absolute)
    if unreliable:
        result["reason"]="；".join(unreliable);result["absolute50_nm"]=None
    else:
        result.update(ic50_nm=float(ic),ci95_nm=ci,status="可估计",reason="通过拟合诊断；仍需人工审核")
    return result

def analyze(raw,record=None,options=None,raw_name="RAW DATA.csv",record_name="记录.xlsx",batch="批次1",import_config=None):
    options=options or {};assay=profile(options)
    from datasets import prepare
    preview=prepare(raw,record,raw_name,record_name,import_config)
    wells=preview['values'];date=preview['date'];meta=preview['metadata'];instrument_background=preview['instrument_background']
    cv_limit=float(options.get("cv_limit",20)); z_limit=float(options.get("z_limit",.5)); span_limit=float(options.get("min_span",20))
    if not 1<=cv_limit<=100 or not 0<=z_limit<1 or not 1<=span_limit<=100:raise ValueError("质控阈值超出允许范围。")
    def rolevals(role):return [wells[d["well"]] for d in meta["design"] if d["role"]==role]
    h=rolevals("no drug"); l=rolevals("cell only"); blank=rolevals("blank")
    if len(h)<2 or len(l)<2:raise ValueError("高信号(high/No Drug)和低信号(low/Cell only)对照至少各需2孔；激活实验请将最大响应对照标high、基础对照标low。")
    high=float(np.mean(h)); low=float(np.mean(l))
    if high<=low:raise ValueError("No Drug/high均值未高于Cell only/low，无法归一化；请检查对照映射。")
    def cv(a):return float(np.std(a,ddof=1)/abs(np.mean(a))*100) if len(a)>1 and np.mean(a)!=0 else None
    z=1-3*(float(np.std(h,ddof=1))+float(np.std(l,ddof=1)))/(high-low)
    alerts=[]
    if z<z_limit:alerts.append({"level":"warning","text":f"Z′={z:.3f}低于阈值{z_limit}；批次需人工审核","source":"No Drug / Cell only设计孔"})
    for label,a in [("No Drug",h),("Cell only",l)]:
        if cv(a) is None:alerts.append({"level":"info","text":f"{label}均值为零，CV无定义；用信号差和Z′判断对照","source":label})
        elif cv(a)>cv_limit:alerts.append({"level":"warning","text":f"{label} CV超过{cv_limit}%","source":label})
    blank_check={"status":"未设置","count":len(blank),"mean":float(np.mean(blank)) if blank else None,"cv_pct":cv(blank) if blank else None}
    if not blank:alerts.append({"level":"warning","text":"未设置独立Blank孔；Cell only仅作为细胞基础信号，独立空白检查无法完成","source":"protocol / Plate Map"})
    else:
        blank_bad=len(blank)<2 or np.mean(blank)>=high or (cv(blank) is not None and cv(blank)>cv_limit)
        blank_check["status"]="需复核" if blank_bad else "通过演示检查"
        alerts.append({"level":"warning" if blank_bad else "info","text":f"独立Blank：{len(blank)}孔，均值{np.mean(blank):g} RLU；{blank_check['status']}；检查复孔数、CV及是否低于No Drug","source":", ".join(d["well"] for d in meta["design"] if d["role"]=="blank")})
    if instrument_background is None:
        alerts.append({"level":"warning","text":"原始仪器文件未找到Background information记录，背景检查无法完成","source":"RAW DATA.csv"})
    else:
        alerts.append({"level":"info","text":f"仪器Background information的Result={instrument_background:g}；未从原始RLU中扣除，保留仪器原始输出","source":"RAW DATA.csv / Background information"})
        if instrument_background<0:alerts.append({"level":"warning","text":"仪器Background Result为负值，需核对仪器背景校正设置","source":"RAW DATA.csv / Background information"})
    missing_meta=[k for k,v in meta.get("cell",{}).items() if v=="未提供"]
    if missing_meta:alerts.append({"level":"info","text":"细胞来源/代次信息缺失，需补充；不从样例报告补写","source":"protocol / 细胞信息"})
    samples=[]
    for name in preview["samples"]:
        design=[d for d in meta["design"] if d.get("sample")==name]
        x=[d["concentration"] for d in design]; y=[wells[d["well"]] for d in design]
        points=[]
        for concentration in sorted(set(x)):
            ds=[d for d in design if d["concentration"]==concentration]; vals=[wells[d["well"]] for d in ds]; c=cv(vals)
            if c is not None and c>cv_limit:alerts.append({"level":"warning","text":f"{name} {concentration:g} nM复孔CV={c:.1f}%；疑似异常复孔，需复核；保留全部孔","source":", ".join(d["well"] for d in ds)})
            if len(vals)<2:alerts.append({"level":"warning","text":f"{name} {concentration:g} nM仅单孔，无法检验复孔离散度","source":ds[0]["well"]})
            points.append({"concentration":concentration,"mean_rlu":float(np.mean(vals)),"sd_rlu":float(np.std(vals,ddof=1)) if len(vals)>1 else 0,"cv_pct":c,"inhibition_pct":float(effect(np.mean(vals),high,low,assay)),"effect_pct":float(effect(np.mean(vals),high,low,assay)),"wells":[{"well":d["well"],"rlu":wells[d["well"]],"source":d["source"]} for d in ds]})
        positive=[(c,v) for c,v in zip(x,y) if c>0]
        if len(positive)<len(x):alerts.append({'level':'info','text':name+'：零浓度观测保留，仅正浓度参与log轴拟合。','source':name})
        fit=fit_sample([p[0] for p in positive],[p[1] for p in positive],high,low,mode='activation' if assay['direction']=='increasing' else 'inhibition',min_span=span_limit)
        if assay['type']=='viability':
            fit['cc50_nm']=fit['ic50_nm'];fit['ic50_nm']=None;fit['reason']=fit['reason'].replace('IC50','CC50')
        fit['potency_nm']=fit[assay['potency_field']];fit['potency_label']=assay['potency_label']
        fit['max_inhibition_pct']=points[-1]['inhibition_pct'];fit['max_observed_inhibition_pct']=max(p['inhibition_pct'] for p in points)
        fit['max_effect_pct']=fit['max_inhibition_pct'];fit['max_observed_effect_pct']=fit['max_observed_inhibition_pct']
        for point in fit['curve']:point['effect_pct']=point['inhibition_pct']
        if fit.get("log_ic50") is not None and fit["residual_sd_rlu"]>0:
            actual_x=np.asarray([p[0] for p in positive]);actual_y=np.asarray([p[1] for p in positive])
            predicted=four_pl(actual_x,fit["bottom"],fit["top"]-fit["bottom"],fit["log_ic50"],-fit["hill"] if assay['direction']=='increasing' else fit['hill'])
            residual=actual_y-predicted
            for d,resid in zip([d for d in design if d['concentration']>0],residual):
                if abs(resid)>3*fit["residual_sd_rlu"]:
                    alerts.append({"level":"warning","text":f'{name} {d["well"]}残差超过3倍残差SD，疑似离群孔；保留数据，需人工复核',"source":d["well"]+" / "+d["source"]})
        if fit["status"]!="可估计":alerts.append({"level":"warning","text":name+"："+fit["reason"],"source":", ".join(d["well"] for d in design)})
        elif z<z_limit or (cv(h) is not None and cv(h)>cv_limit) or (cv(l) is not None and cv(l)>cv_limit):fit["status"]="需复核（批次质控未通过）"
        display=clean(options.get("sample_names",{}).get(name,name))
        samples.append({"sample":name,"display_name":display,"points":points,**fit,"sample_metadata":{"manufacturer":"未提供","lot":"未提供","stock_concentration":"未提供","molecular_weight":"未提供"}})
    ref_name=str(options.get("reference_sample","")).strip()
    reference={"sample":ref_name or None,"status":"未指定","ic50_nm":None,"expected_ic50_nm":None,"ratio_to_expected":None}
    expected=number(options.get("reference_ic50"))
    if ref_name:
        ref=next((s for s in samples if s["sample"]==ref_name),None)
        if ref is None:raise ValueError("参考内控必须是实验设计中的样品，不能使用不存在的名称。")
        if options.get("reference_ic50") not in (None,"") and (expected is None or expected<=0):raise ValueError("参考内控预期IC50需为正数（nM）。")
        reference.update(ic50_nm=potency(ref),potency_nm=potency(ref),potency_label=assay['potency_label'],expected_ic50_nm=expected,status="预期效价未提供")
        if potency(ref) is None or ref["status"]!="可估计":reference["status"]="需复核：参考内控拟合或对照未通过"
        elif expected is not None:
            ratio=potency(ref)/expected
            reference.update(ratio_to_expected=ratio,status="通过演示检查" if .5<=ratio<=2 else "需复核：与预期效价差异超过2倍")
        alerts.append({"level":"info" if reference["status"]=="通过演示检查" else "warning","text":f"参考内控{ref_name}：{reference['status']}；由用户指定，预期值不是从Report.docx读取","source":ref_name+" / 用户确认"})
        if reference["status"].startswith("需复核"):
            for sample in samples:
                if sample["status"]=="可估计":
                    sample["status"]="需复核（参考内控未通过）"
    else:
        alerts.append({"level":"warning","text":"未指定药效参考内控及预期IC50，参考药物效价检查无法完成；不将任一Drug自动当作内控","source":"protocol / 用户输入"})
    unused=sorted(set(wells)-{d["well"] for d in meta["design"]})
    if unused:alerts.append({"level":"info","text":"未关联设计的原始孔未参与计算："+", ".join(unused),"source":"CSV / Plate Map"})
    passed=z>=z_limit and all(cv(a) is None or cv(a)<=cv_limit for a in (h,l))
    rules={**RULES,"background":"仪器Background Result仅读取核查，不重复扣背景；独立Blank检查孔数、CV与是否低于No Drug；缺失时不宣称通过","reference":"药效参考内控由用户确认其样品身份与预期IC50；观测/预期0.5–2倍是演示接受区间，需替换为业务SOP"}
    if assay['type']!='inhibition':
        for alert in alerts:alert['text']=alert['text'].replace('IC50',assay['potency_label']).replace('Cell only仅作为细胞基础信号','low仅作为低信号对照').replace('RLU',assay['readout'])
        rules['normalization']='效应率 = 100 × (读数 − 低信号均值) / (高信号均值 − 低信号均值)，不截断' if assay['direction']=='increasing' else '效应率 = 100 × (高信号均值 − 读数) / (高信号均值 − 低信号均值)，不截断'
        rules['fit']='原始信号全复孔等权重4PL；Bottom+(Top-Bottom)/(1+10^((log10(C)-logMidpoint)*'+('-Hill' if assay['direction']=='increasing' else 'Hill')+'))'
        rules['ic50']='相对'+assay['potency_label']+'为拟合上下平台中点；绝对50%效应浓度另列；不外推临床效力。'
        rules['reference']=rules['reference'].replace('IC50',assay['potency_label'])
    rules['assay']=assay['principle']
    rules['import']='所有终浓度统一为nM；零浓度仅保留观测，不进入对数拟合。'
    return dict(batch=batch,date=date,assay=assay,metadata=meta,samples=samples,controls={"no_drug_mean":high,"cell_only_mean":low,"no_drug_cv":cv(h),"cell_only_cv":cv(l),"z_prime":float(z),"blank_mean":float(np.mean(blank)) if blank else None,"blank_check":blank_check,"reference_control":reference,"instrument_background":instrument_background,"passed":passed,"scope":"passed仅表示高/低信号窗口与CV，不代表所有质控项目通过"},alerts=alerts,options={"cv_limit":cv_limit,"z_limit":z_limit,"min_span":span_limit,"reference_sample":ref_name,"reference_ic50":expected},rules=rules,provenance={"raw_file":raw_name,"record_file":record_name if record else '单文件集成设计',"raw_sha256":preview['raw_sha256'],"record_sha256":preview['record_sha256'],"raw_wells":wells,"design":meta["design"],"import":preview['import_info']})
