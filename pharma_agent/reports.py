import io,json,html,threading,os
from pathlib import Path
os.environ.setdefault('MPLCONFIGDIR',str(Path(__file__).resolve().parent/'.local'/'matplotlib'))
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches,Pt,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from assays import potency
PLOT_LOCK=threading.Lock()
COLORS=["#0c8c7b","#6485e8","#ed9c41","#cc6b91"]

def chart_png(results):
    with PLOT_LOCK:
        fig,axes=plt.subplots(1,len(results),figsize=(7*len(results),3.1),squeeze=False)
        for ax,r in zip(axes[0],results):
            for i,s in enumerate(r["samples"]):
                color=COLORS[i%4];p=s["points"]
                p=[v for v in p if v['concentration']>0]
                xs=[v["concentration"] for v in p];ys=[v["inhibition_pct"] for v in p]
                scale=r["controls"]["no_drug_mean"]-r["controls"]["cell_only_mean"]
                err=[v["sd_rlu"]/scale*100 for v in p]
                ax.errorbar(xs,ys,yerr=err,fmt="o",markersize=4,color=color,label=s["sample"],capsize=2)
                if s["curve"]:ax.plot([v["concentration"] for v in s["curve"]],[v["inhibition_pct"] for v in s["curve"]],color=color)
            ax.set_xscale("log");ax.axhline(50,color="#a5aebb",ls="--",lw=.8)
            ax.set_xlabel("Final concentration (nM)");ax.set_ylabel("Activation (%)" if r.get('assay',{}).get('direction')=='increasing' else 'Effect (%)');ax.set_title("Batch "+str(results.index(r)+1));ax.legend(fontsize=8);ax.grid(alpha=.15)
        fig.tight_layout();out=io.BytesIO();fig.savefig(out,format="png",dpi=160);plt.close(fig);return out.getvalue()
def fmt(n,d=4):return "不可估计" if n is None else f"{n:.{d}g}"
def summary(r):
    a=r.get('assay',{});label=a.get('potency_label','IC50')
    valid=[s for s in r["samples"] if potency(s) is not None]
    line="；".join(f'{s["display_name"]}相对{label} {fmt(potency(s))} nM' for s in valid)
    bad=[s["display_name"] for s in r["samples"] if potency(s) is None]
    purpose=r["metadata"]["purpose"].replace("实验目的：","").rstrip("。")
    return f'目的：{purpose}；方法：依据确认后的终浓度与观测，对{a.get("readout","RLU")}进行复孔/对照质控、四参数拟合和效应率归一化。{r["batch"]}结果：{line or "无可靠效价"}。'+("、".join(bad)+"不可可靠估计。" if bad else "")+"结论：仅反映当前实验条件下的体外浓度响应，需人工审核。"
def make_brief_docx(result,index=0,edited_summary=None):
    r=result["results"][index];m=r["metadata"];a=r.get('assay',{});label=a.get('potency_label','IC50');readout=a.get('readout','RLU');doc=Document();sec=doc.sections[0]
    sec.page_height=Cm(29.7);sec.page_width=Cm(21);sec.top_margin=Cm(1.2);sec.bottom_margin=Cm(1.1);sec.left_margin=Cm(1.5);sec.right_margin=Cm(1.5)
    style=doc.styles["Normal"];style.font.name="Microsoft YaHei";style.font.size=Pt(8)
    style._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
    style.paragraph_format.space_after=Pt(3)
    style.paragraph_format.line_spacing=1.1
    def p(text, bold=False, size=None, center=False):
        para=doc.add_paragraph(style="Normal")
        para.paragraph_format.space_after=Pt(3)
        if center:para.alignment=1
        run=para.add_run(str(text));run.bold=bold;run.font.name="Microsoft YaHei";run._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
        if size:run.font.size=Pt(size)
        return para
    def heading(text):
        return p(text,bold=True,size=10)
    def table(headers,rows):
        tb=doc.add_table(rows=1,cols=len(headers))
        tb.style="Light Shading Accent 1"
        for c,v in zip(tb.rows[0].cells,headers):c.text=str(v)
        for row in rows:
            for c,v in zip(tb.add_row().cells,row):c.text=str(v)
        for row in tb.rows:
            for c in row.cells:
                for para in c.paragraphs:
                    para.paragraph_format.space_after=Pt(0)
                    para.paragraph_format.space_before=Pt(0)
                    for run in para.runs:
                        run.font.size=Pt(7);run.font.name="Microsoft YaHei"
                        run._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
        return tb
    p(a.get('label','体外药理')+"浓度响应实验报告",bold=True,size=16,center=True)
    p(f'{r["batch"]} · 测量时间：{r["date"]} · 报告初稿 / 待人工审核',size=8)
    heading("1. 审批表")
    table(["角色","部门","姓名 / 签名 / 日期"],[["起草","体外药理部","待填写"],["审核","质量管理部","待填写"],["批准","体外药理部","待填写"]])
    heading("2. 摘要")
    p(edited_summary or summary(r))
    heading("3. 缩略语")
    table(["英文简称","英文全称","中文全称"],[
        ["RLU","Relative Light Unit","相对光单位"],
        ["IC50","Half maximal inhibitory concentration","半数抑制浓度"],
        ["EC50","Half maximal effective concentration","半数效应浓度（激活分析使用）"],
        ["CC50","Half maximal cytotoxic concentration","本项目采用活力下降曲线相对中点，不等于已证实细胞死亡"],
        ["nM","nmol/L","纳摩尔每升"],
        ["μM","μmol/L","微摩尔每升"],
    ])
    heading("4. 实验目的、仪器、试剂与细胞")
    p(m["purpose"])
    p("仪器："+ "；".join(f"{a[0]} / {a[1]} / {a[2]}" for a in m["instruments"])+"。设备编号未提供。")
    p("试剂耗材："+ "；".join(f"{a[0]} / {a[1]} / {a[2]}" for a in m["reagents"])+"。批号未提供。")
    p("细胞："+m.get("cell",{}).get("name","未提供")+"；来源、代次："+m.get("cell",{}).get("source","未提供")+" / "+m.get("cell",{}).get("passage","未提供"))
    heading("5. 检测对象")
    table(["样品","厂家 / 批号","库存浓度 / 分子量","测试终浓度"],[[s["display_name"],"未提供 / 未提供","未提供 / 未提供",f'{min(v["concentration"] for v in s["points"]):g}–{max(v["concentration"] for v in s["points"]):g} nM'] for s in r["samples"]])
    heading("6. 实验原理")
    p(a.get('principle','按高低信号对照评价浓度响应。')+f'读数：{readout}。高/低信号对照由用户显式确认，具体靶点、机制与湿实验资料不自动推断。独立Blank状态：'+r['controls']['blank_check']['status']+'。')
    heading("7. 实验方法")
    methods=[a["text"] for a in m["methods"] if "浓度范围" not in a["text"] and "数据处理" not in a["text"]]
    cs=sorted({v["concentration"] for s in r["samples"] for v in s["points"]},reverse=True)
    for text in [
        "1）细胞收集与铺板："+" ".join(t for t in methods if "收集细胞" in t or "细胞铺板" in t),
        "2）药物配制与加板：终浓度为"+", ".join(f"{v:g}" for v in cs)+" nM；"+" ".join(t for t in methods if "铺板体系" in t),
        "3）孵育："+" ".join(t for t in methods if t.startswith("孵育"))+"；时长未提供。",
        "4）检测："+" ".join(t for t in methods if "到点" in t or "检测" in t and "试剂" in t or "按照说明书" in t),
        "5）数据处理："+r["rules"]["fit"]+"；"+r["rules"]["normalization"]+"；异常仅标记不删孔。"]:
        p(text)
    doc.add_page_break()
    heading("8. 实验结果")
    doc.add_picture(io.BytesIO(chart_png([r])),width=Cm(17.5))
    p("点为复孔均值，误差线为SD；曲线由输入信号的4PL拟合获得；零浓度不进入对数图，原始记录保留。合成示例不代表真实实验。")
    table(["样品","相对"+label+" / nM","近似95%区间 / nM","最高浓度效应率","R² / 状态"],[[s["display_name"],fmt(potency(s)),("–".join(fmt(x) for x in s["ci95_nm"]) if s["ci95_nm"] else "不可估计"),fmt(s["max_inhibition_pct"])+"%",fmt(s["r2"])+" / "+s["status"]] for s in r["samples"]])
    p("测试范围最大观测"+a.get('effect_label','效应率')+"："+"；".join(f'{s["display_name"]} {s["max_observed_inhibition_pct"]:.1f}%' for s in r["samples"])+"。不代表已证明的生物学最大效应；归一化超过100%或小于0%时保留原值。")
    c=r["controls"];p(f'对照：高信号 {c["no_drug_mean"]:.4g} {readout}（CV {fmt(c["no_drug_cv"])}%）；低信号 {c["cell_only_mean"]:.4g} {readout}（CV {fmt(c["cell_only_cv"])}%）；Z′={c["z_prime"]:.3f}；仪器Background Result={fmt(c.get("instrument_background"))}（未扣除）。')
    p("质控 / 拟合提示："+"；".join(a["text"] for a in r["alerts"]))
    p("相对"+label+"对应拟合上下平台中点，并非绝对50%效应。绝对50%效应浓度："+ "；".join(s["display_name"]+" "+fmt(s["absolute50_nm"])+" nM" for s in r["samples"])+ "。近似区间基于log中点协方差与t分布，不等同Prism的profile-likelihood区间。")
    heading("9. 附件、溯源与审核记录")
    pr=r["provenance"];p("输入："+pr["raw_file"]+" + "+pr["record_file"]+"。原始长表/矩阵与显式设计参与计算；Result工作表未读取为分析输入。")
    p("SHA256 RAW："+pr["raw_sha256"]+"\nSHA256记录："+(pr["record_sha256"] or '单文件模式，无第二文件'))
    p('导入配置：'+json.dumps(pr.get('import',{}),ensure_ascii=False))
    p("处理规则："+r["rules"]["qc"]+f'。当前CV阈值 {r["options"]["cv_limit"]}%、Z′阈值 {r["options"]["z_limit"]}、最低动态范围 {r["options"]["min_span"]}%。')
    p("批次比较：仅作描述性比较；需确认样品身份、条件与内控一致。孔位、原始RLU、设计单元格、参数和工具调用见随附审计JSON。")
    if len(result["results"])>1:
        p("本次各批次相对效价（nM，仅同场景比较）："+"；".join(b["batch"]+"："+", ".join(s["display_name"]+" "+fmt(potency(s)) for s in b["samples"]) for b in result["results"]))
    p("内控检查："+r["controls"].get("reference_control",{}).get("status","未指定")+"；独立Blank："+r["controls"].get("blank_check",{}).get("status","未设置")+"。")
    p("AI与人工记录：报告初稿由确定性工具生成；解读模式："+result["mode"]+"。"+("摘要由用户编辑，未核验修改内容。" if edited_summary else "用户尚未编辑摘要。")+"起草/审核/批准人员与签名均未填写，需人工核对。")
    p("核验清单：①确认Plate Map与终浓度及对照；②复核拟合状态、Drug4与缺失信息；当前均待签署。")
    footer=sec.footer.paragraphs[0];footer.text="AssayPilot · 演示报告 / 待人工审核 · "
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");footer._p.append(field)
    out=io.BytesIO();doc.save(out);return out.getvalue()
def make_docx(result,index=0,edited_summary=None):
    """Self-authored generic report: no external DOCX, artwork or enterprise data."""
    doc=Document(io.BytesIO(make_brief_docx(result,index,edited_summary)))
    r=result["results"][index];sec=doc.sections[0]
    sec.top_margin=Cm(2.9);sec.header_distance=Cm(.6)
    header=sec.header.add_table(rows=2,cols=3,width=Cm(18))
    header.style="Table Grid"
    for cell,value in zip(header.rows[0].cells,["AssayPilot", "通用实验报告 · 待审核", "SYNTHETIC DEMO" if "合成" in r["batch"] else "用户上传数据"]):cell.text=value
    for cell,value in zip(header.rows[1].cells,[r["batch"],"仅为报告初稿",""]):cell.text=value
    page=header.cell(1,2).paragraphs[0];page.add_run("第 ")
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");page._p.append(field);page.add_run(" 页")
    for row in header.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:run.font.size=Pt(8)
    # Uploaded descriptions and interpretations remain data, never template instructions.
    doc.core_properties.title="AssayPilot generic assay report"
    doc.core_properties.author="AssayPilot contributors"
    out=io.BytesIO();doc.save(out);return out.getvalue()

def make_html(result,index=0,edited_summary=None):
    import base64
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc=Document(io.BytesIO(make_brief_docx(result,index,edited_summary)))
    blocks=['<section class="sheet">']
    for el in doc.element.body:
        if el.tag==qn("w:p"):
            p=Paragraph(el,doc)
            if any(b.get(qn("w:type"))=="page" for b in el.iter(qn("w:br"))):
                blocks.append('</section><section class="sheet">')
                continue
            drawings=list(el.iter(qn("a:blip")))
            if drawings:
                for drawing in drawings:
                    part=doc.part.related_parts[drawing.get(qn("r:embed"))]
                    blocks.append('<img alt="剂量响应曲线" src="data:image/png;base64,'+base64.b64encode(part.blob).decode()+'">')
            elif p.text:
                tag="h1" if p.style.name=="Title" or any(run.font.size and run.font.size.pt==16 for run in p.runs) else "h2" if p.style.name.startswith("Heading") or any(run.bold and run.font.size and run.font.size.pt==10 for run in p.runs) else "p"
                blocks.append("<"+tag+">"+html.escape(p.text).replace("\n","<br>")+"</"+tag+">")
        elif el.tag==qn("w:tbl"):
            table=Table(el,doc);blocks.append("<table>")
            for i,row in enumerate(table.rows):
                tag="th" if i==0 else "td"
                blocks.append("<tr>"+"".join("<"+tag+">"+html.escape(c.text)+"</"+tag+">" for c in row.cells)+"</tr>")
            blocks.append("</table>")
    blocks.append("</section>")
    css="*{box-sizing:border-box}body{margin:0;background:#edf1f4;color:#172c39;font:8pt 'Microsoft YaHei',sans-serif;line-height:1.45}.sheet{width:210mm;min-height:297mm;margin:18px auto;padding:12mm 15mm 11mm;background:white;box-shadow:0 1px 8px #0001}h1{font-size:16pt;margin:0 0 5pt}h2{font-size:10pt;color:#087766;margin:5pt 0 3pt}p{margin:0 0 3pt;overflow-wrap:anywhere}table{width:100%;border-collapse:collapse;font-size:7pt;margin-bottom:4pt}th,td{border:1px solid #dce5ea;text-align:left;padding:3pt 4pt}th{background:#eaf5f1}img{width:175mm;height:auto;display:block;margin:5pt 0}@page{size:A4;margin:0}@media print{body{background:white}.sheet{min-height:0;margin:0;box-shadow:none;break-after:page}.sheet:last-child{break-after:auto}}"
    return ('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>标准实验报告</title><style>'+css+'</style>'+"".join(blocks)+"</html>").encode()
