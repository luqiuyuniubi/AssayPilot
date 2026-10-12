"""Bounded format adapters; all layouts become explicit measurements + design."""
import csv,io,json,math,re,zipfile
from pathlib import Path
from openpyxl import load_workbook
from defusedxml.ElementTree import iterparse

EXTENSIONS={'.csv','.tsv','.txt','.xlsx','.json'}
ALIASES={
 'well':['well','wellid','wellposition','position','孔位','孔号'],
 'sample':['sample','sampleid','compound','compoundid','drug','样品','样品名称','化合物','药物'],
 'concentration':['concentration','conc','dose','浓度','终浓度','剂量'],
 'value':['value','response','signal','rlu','rfu','od','absorbance','luminescence','viability','读数','响应','信号','吸光度','存活率'],
 'role':['role','control','type','孔类型','对照','角色'],
 'unit':['unit','concentrationunit','单位','浓度单位'],
}
FACTORS={'pm':.001,'nm':1.,'um':1000.,'mm':1e6,'m':1e9}
def key(s):return re.sub(r'[\s_\-()%/\[\]]','',str(s).strip().lower().replace('μ','u').replace('µ','u'))
def finite(value,label):
    try:n=float(value)
    except (ValueError,TypeError):raise ValueError(label+'必须为有效数字。') from None
    if not math.isfinite(n):raise ValueError(label+'不能是NaN或无穷大。')
    return n
def unit_factor(unit):
    name=str(unit).strip().replace('μ','u').replace('µ','u').lower()
    if name not in FACTORS:raise ValueError('浓度单位不支持：'+str(unit)+'；支持pM/nM/µM/mM/M，不自动换算质量浓度。')
    return FACTORS[name]
def check_xlsx(data):
    if len(data)>8*1024*1024:raise ValueError('实验记录超过8MB限制。')
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        entries=z.infolist()
        if len(entries)>2000 or sum(i.file_size for i in entries)>40*1024*1024:raise ValueError('实验记录解压体积或条目数过大。')
        cells=0
        for item in entries:
            if item.file_size>12*1024*1024:raise ValueError('实验记录单个压缩条目过大。')
            if not item.filename.endswith(('.xml','.rels')):continue
            sheet=item.filename.startswith('xl/worksheets/')
            with z.open(item) as source:
                for event,e in iterparse(source,events=('start','end'),forbid_dtd=True):
                    if event=='end':e.clear();continue
                    if not sheet:continue
                    if e.tag.startswith('{') and not e.tag.startswith('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'):continue
                    tag=e.tag.rsplit('}',1)[-1]
                    if tag=='row' and 'r' in e.attrib and (not str(e.get('r')).isdigit() or not 1<=int(e.get('r'))<=5000):raise ValueError('工作表行范围超过安全限制。')
                    if tag=='col' and ('min' in e.attrib or 'max' in e.attrib) and any(not str(e.get(k,'')).isdigit() or not 1<=int(e.get(k))<=256 for k in ('min','max')):raise ValueError('工作表列范围超过安全限制。')
                    if tag=='c':
                        cells+=1
                        if cells>150000:raise ValueError('实验记录单元格数量超过安全限制。')
                    if tag in ('c','dimension','mergeCell'):
                        for column,row in re.findall(r'([A-Z]+)(\d+)',e.get('r') or e.get('ref') or ''):
                            col=0
                            for letter in column:col=col*26+ord(letter)-64
                            if not 1<=int(row)<=5000 or col>256:raise ValueError('实验记录工作表范围超过安全限制。')
def read_tables(data,filename):
    if not data or len(data)>8*1024*1024:raise ValueError('文件需为非空且不超过8MB。')
    ext=Path(filename).suffix.lower()
    if ext not in EXTENSIONS:raise ValueError('数值文件支持CSV、TSV、TXT、XLSX、JSON；不从图片或PDF猜测数值。')
    metadata={}
    if ext=='.xlsx':
        check_xlsx(data);wb=load_workbook(io.BytesIO(data),data_only=True,read_only=True,keep_links=False)
        try:
            tables={ws.title:[[v for v in row] for row in ws.iter_rows(values_only=True)] for ws in wb if ws.title.lower() not in ('result','results','结果')}
        finally:wb.close()
    elif ext=='.json':
        try:doc=json.loads(data.decode('utf-8-sig'))
        except (ValueError,UnicodeError,RecursionError):raise ValueError('JSON内容无效或嵌套过深。') from None
        if isinstance(doc,dict):
            metadata=doc.get('metadata',{});records=doc.get('measurements',doc.get('rows',doc.get('data')))
            if not isinstance(metadata,dict):raise ValueError('JSON metadata需为对象。')
        else:records=doc
        if not isinstance(records,list) or not records or len(records)>5000 or not all(isinstance(r,dict) for r in records):raise ValueError('JSON需为最多5000行的对象数组，或包含measurements数组的对象。')
        columns=list(dict.fromkeys(k for r in records for k in r))
        tables={'data':[columns]+[[r.get(k) for k in columns] for r in records]}
    else:
        text=None
        for encoding in ('utf-8-sig','gb18030'):
            try:text=data.decode(encoding);break
            except UnicodeError:pass
        if text is None:raise ValueError('文本编码需为UTF-8或GB18030。')
        try:delimiter='\t' if ext=='.tsv' else csv.Sniffer().sniff(text[:8192],delimiters=',\t;').delimiter
        except csv.Error:delimiter=','
        rows=[]
        for row in csv.reader(io.StringIO(text),delimiter=delimiter):
            if len(rows)>=5000:raise ValueError('表格行数超过5000。')
            if len(row)>256:raise ValueError('表格列数超过256。')
            rows.append(row)
        tables={'data':rows}
    if not tables:raise ValueError('没有可读取工作表；Result工作表不是数值输入。')
    for rows in tables.values():
        if len(rows)>5000 or any(len(r)>256 for r in rows):raise ValueError('表格范围超出安全限制。')
    return tables,metadata
def mapping_for(columns,manual=None):
    manual=manual or {};mapping={}
    if not isinstance(manual,dict):raise ValueError('列映射必须为对象。')
    for field,aliases in ALIASES.items():
        if field in manual and manual[field]:
            if manual[field] not in columns:raise ValueError('映射列不存在：'+str(manual[field]))
            mapping[field]=columns.index(manual[field]);continue
        matches=[]
        for i,column in enumerate(columns):
            k=key(column)
            base=re.sub(r'(?:pm|nm|um|mm)$','',k) if field=='concentration' else k
            if k in aliases or base in aliases:matches.append(i)
        if len(matches)>1:raise ValueError('多个列可映射为'+field+'，请在页面选择正确列。')
        if matches:mapping[field]=matches[0]
    return mapping
def select_table(tables,config,needs):
    sheet=config.get('sheet');manual=config.get('columns',{});header=config.get('header_row')
    if sheet and sheet not in tables:raise ValueError('工作表不存在：'+str(sheet))
    candidates=[]
    for name,rows in tables.items():
        if sheet and sheet!=name:continue
        positions=[int(header)-1] if header else range(min(30,len(rows)))
        for pos in positions:
            if pos<0 or pos>=len(rows):raise ValueError('表头行超出范围。')
            columns=[str(v).strip() if v is not None else '' for v in rows[pos]]
            try:m=mapping_for(columns,manual)
            except ValueError:
                if sheet and header:raise
                continue
            if all(n in m for n in needs):candidates.append((name,rows,pos,columns,m));break
    if not candidates:raise ValueError('未识别表头。请选择工作表、表头行并映射列；长表至少需要读数，设计还需样品和浓度。')
    if len(candidates)>1 and not sheet:raise ValueError('多个工作表符合要求，请显式选择要读取的工作表。')
    return candidates[0]
def describe(data,filename):
    tables,_=read_tables(data,filename);out=[]
    for name,rows in tables.items():
        candidates=[]
        for i,row in enumerate(rows[:30]):
            if not any(v is not None and str(v).strip() for v in row):continue
            cols=[str(v or '') for v in row]
            try:m=mapping_for(cols)
            except ValueError:m={}
            # A data row can contain the literal word "sample"; do not confuse
            # this with a useful header. Arbitrary headers remain selectable.
            if len(m)<2 or any(isinstance(v,(int,float)) for v in row):m={}
            candidates.append({'header_row':i+1,'columns':cols,'suggested':{k:cols[v] for k,v in m.items()}})
        candidates.sort(key=lambda h:(-len(h['suggested']),-len([c for c in h['columns'] if c]),h['header_row']))
        out.append({'name':name,'headers':candidates,'row_count':len(rows)})
    return {'filename':filename,'tables':out}
def role(value):
    k=key(value)
    if not k or k in ('sample','test','样品','样本'):return 'sample'
    if k in ('nodrug','high','highsignal','vehicle','vehiclecontrol','高信号','溶剂对照','未加药'):return 'no drug'
    if k in ('cellonly','low','lowsignal','baseline','低信号','基础对照'):return 'cell only'
    if k in ('blank','空白'):return 'blank'
    raise ValueError('未知对照类型：'+str(value)+'；请用sample/high/low/blank，勿猜测对照含义。')
def well(value):
    token=str(value or '').strip().upper()
    m=re.fullmatch(r'([A-P])0*(\d{1,2})',token)
    if not m or not 1<=int(m[2])<=24:raise ValueError('孔位需为A1–P24；非孔板长表可省略孔位，由行号标识。')
    return m[1]+str(int(m[2]))
def make_metadata(design,provided=None):
    provided=provided or {}
    return {'title':str(provided.get('title','体外药理浓度响应实验'))[:250],'purpose':str(provided.get('purpose','实验目的：评价当前条件下的浓度响应关系'))[:2000],
        'cell':{'name':str(provided.get('cell','未提供'))[:250],'source':'未提供','passage':'未提供'},'instruments':[],'reagents':[],
        'methods':[{'text':'按长表的样品、终浓度和读数进行质控及4PL分析；原始观测保留，具体湿实验方法待补充。','source':'用户表格'}],
        'design':design,'scope':'输入长表与用户确认；单位统一至nM，Result工作表不读取。'}
def long_table(data,filename,config=None,require_value=True,require_design=False):
    config=config or {};tables,metadata=read_tables(data,filename)
    needs=['sample','concentration'] if require_design else ['value']
    if require_design and require_value:needs.append('value')
    sheet,rows,header,columns,m=select_table(tables,config,needs)
    if not require_design and not all(x in m for x in ('well','sample','concentration')) and 'well' not in m and not ('sample' in m and 'concentration' in m):raise ValueError('仅读数表需提供孔位；单文件长表需样品和浓度。')
    design=[];values={};conversions=set();seen=set()
    def value(row,field):return row[m[field]] if field in m and m[field]<len(row) else None
    embedded='sample' in m and 'concentration' in m
    for offset,row in enumerate(rows[header+1:],header+2):
        if not any(v is not None and str(v).strip() for v in row):continue
        identifier=well(value(row,'well')) if 'well' in m else 'row-'+str(offset)
        if identifier in seen:raise ValueError('孔位重复：'+identifier+'；请把多块板分为不同批次。')
        seen.add(identifier)
        if require_value:
            signal=finite(value(row,'value'),f'{sheet}:{offset}读数')
            if signal<0:raise ValueError('当前按非负原始信号分析；请提供原始读数和高低信号对照，不从负百分比猜测原值。')
            values[identifier]=signal
        if embedded:
            r=role(value(row,'role'));sample=value(row,'sample')
            # Control labels may be carried in the sample column only when there
            # is no explicit role column; otherwise user roles take precedence.
            if r=='sample' and 'role' not in m:
                try:r=role(sample)
                except ValueError:r='sample'
            entry={'well':identifier,'role':r,'source':f'{filename} / {sheet}!row{offset}'}
            if r=='sample':
                if sample is None or not str(sample).strip():raise ValueError('样品名称不能为空。')
                label=columns[m['concentration']];match=re.search(r'(pM|nM|[uμµ]M|mM|M)\s*[)\]]?$',label)
                declared=value(row,'unit') or config.get('unit') or (match[1] if match else 'nM')
                if match and unit_factor(declared)!=unit_factor(match[1]):raise ValueError('单位列/选择单位与浓度表头冲突，请修正而非静默覆盖。')
                concentration=finite(value(row,'concentration'),'终浓度')*unit_factor(declared)
                if concentration<0:raise ValueError('终浓度不能为负数。')
                conversions.add(str(declared)+'→nM')
                entry.update(sample=str(sample).strip()[:100],concentration=concentration,unit='nM',original_concentration=value(row,'concentration'),original_unit=str(declared))
            design.append(entry)
    if not values and require_value:raise ValueError('没有有效读数。')
    if require_design and not design:raise ValueError('没有有效样品设计。')
    return {'values':values,'design':design,'metadata':make_metadata(design,metadata),'mapping':{field:columns[idx] for field,idx in m.items()},'sheet':sheet,'header_row':header+1,'conversions':sorted(conversions)}
def matrix(data,filename,config=None):
    tables,_=read_tables(data,filename);config=config or {};sheet=config.get('sheet')
    if sheet and sheet not in tables:raise ValueError('矩阵工作表不存在。')
    candidates=[]
    for name,rows in tables.items():
        if sheet and name!=sheet:continue
        wells={};active=None
        for row in rows:
            numbered=[str(v).strip() for v in row[1:]]
            if len(numbered)>=12 and numbered[:12]==[str(i) for i in range(1,13)]:
                n=24 if numbered[:24]==[str(i) for i in range(1,25)] else 12
                if active is not None:raise ValueError('一份矩阵含多块板，请分批上传。')
                active=n;continue
            if active and row and str(row[0]).strip().upper() in list('ABCDEFGHIJKLMNOP'):
                label=str(row[0]).strip().upper()
                if active==12 and label not in 'ABCDEFGH':raise ValueError('96孔矩阵行只能为A–H。')
                for col,v in enumerate(row[1:active+1],1):
                    if v is None or not str(v).strip():continue
                    identifier=label+str(col)
                    if identifier in wells:raise ValueError('重复孔位：'+identifier)
                    n=finite(v,identifier)
                    if n<0:raise ValueError('原始信号不能为负。')
                    wells[identifier]=n
        if wells:candidates.append((name,wells))
    if len(candidates)!=1:raise ValueError('未找到唯一96/384孔矩阵，请选工作表或改用长表。')
    return candidates[0][1],candidates[0][0]
